from fastapi import APIRouter, Form
from fastapi.responses import StreamingResponse
import io
import tempfile
import os
from datetime import datetime
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re

router = APIRouter()


def register_chinese_font():
    font_paths = [
        "/usr/share/fonts/truetype/wqy/wqy-microhei.ttc",
        "C:/Windows/Fonts/simhei.ttf"
    ]
    for path in font_paths:
        if os.path.exists(path):
            try:
                pdfmetrics.registerFont(TTFont('ChineseFont', path))
                return 'ChineseFont'
            except:
                continue
    return 'Helvetica'


FONT_NAME = register_chinese_font()


@router.post("/generate_report")
async def generate_report(report_type: str = Form(...), content: str = Form(...)):
    if report_type == "docx":
        doc = Document()
        title = doc.add_heading('心镜 - 心理健康报告', 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph(f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")

        lines = content.split('\n')
        for line in lines:
            if line.strip() == "":
                continue
            if line.startswith('# '):
                doc.add_heading(line[2:], level=1)
                continue
            elif line.startswith('## '):
                doc.add_heading(line[3:], level=2)
                continue

            parts = re.split(r'(\*\*.*?\*\*)', line)
            if parts:
                p = doc.add_paragraph()
                for part in parts:
                    if part.startswith('**') and part.endswith('**'):
                        run = p.add_run(part[2:-2])
                        run.bold = True
                    else:
                        p.add_run(part)

        temp_file = tempfile.NamedTemporaryFile(delete=False, suffix='.docx')
        doc.save(temp_file.name)
        temp_file.close()

        with open(temp_file.name, 'rb') as f:
            content_bytes = f.read()
        os.unlink(temp_file.name)

        return StreamingResponse(
            io.BytesIO(content_bytes),
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": "attachment; filename=mindmirror_report.docx"}
        )
    else:
        temp_file = tempfile.NamedTemporaryFile(delete=False, suffix='.pdf')
        c = canvas.Canvas(temp_file.name, pagesize=A4)
        width, height = A4
        y = height - 50

        c.setFont(FONT_NAME, 20)
        c.drawString(50, y, "心镜 - 心理健康报告")
        y -= 40

        c.setFont(FONT_NAME, 10)
        c.drawString(50, y, f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        y -= 40

        lines = content.split('\n')
        for line in lines:
            if y < 80:
                c.showPage()
                y = height - 50
                c.setFont(FONT_NAME, 11)

            if line.strip() == "":
                y -= 10
                continue

            if line.startswith('# '):
                c.setFont(FONT_NAME, 16)
                c.drawString(50, y, line[2:])
                y -= 25
                c.setFont(FONT_NAME, 11)
                continue
            elif line.startswith('## '):
                c.setFont(FONT_NAME, 14)
                c.drawString(50, y, line[3:])
                y -= 20
                c.setFont(FONT_NAME, 11)
                continue

            clean_line = re.sub(r'\*\*(.*?)\*\*', r'\1', line)
            c.setFont(FONT_NAME, 11)
            c.drawString(50, y, clean_line)
            y -= 20

        c.save()
        temp_file.close()

        with open(temp_file.name, 'rb') as f:
            content_bytes = f.read()
        os.unlink(temp_file.name)

        return StreamingResponse(
            io.BytesIO(content_bytes),
            media_type="application/pdf",
            headers={"Content-Disposition": "attachment; filename=mindmirror_report.pdf"}
        )


@router.get("/export_report")
async def export_report():
    from utils.storage import load_chat_history, load_emotion_history, load_survey_data
    from services.ai_service import call_ai

    chat_history = load_chat_history()
    emotion_history = load_emotion_history()
    survey_data = load_survey_data()

    temp_file = tempfile.NamedTemporaryFile(delete=False, suffix='.pdf')
    c = canvas.Canvas(temp_file.name, pagesize=A4)
    width, height = A4
    y = height - 50

    c.setFont(FONT_NAME, 18)
    c.drawString(50, y, "心镜 - 心理健康报告")
    y -= 40

    c.setFont(FONT_NAME, 10)
    c.drawString(50, y, f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    y -= 40

    c.setFont(FONT_NAME, 14)
    c.drawString(50, y, "一、情绪统计")
    y -= 25
    c.setFont(FONT_NAME, 11)

    emotion_counts = {}
    for item in emotion_history:
        emo = item.get('emotion', '未知')
        emotion_counts[emo] = emotion_counts.get(emo, 0) + 1

    if emotion_counts:
        for emo, count in emotion_counts.items():
            c.drawString(60, y, f"{emo}：{count} 次")
            y -= 20
            if y < 100:
                c.showPage()
                y = height - 50
    else:
        c.drawString(60, y, "暂无情绪记录")
        y -= 20

    y -= 15

    c.setFont(FONT_NAME, 14)
    c.drawString(50, y, "二、对话统计")
    y -= 25
    c.setFont(FONT_NAME, 11)
    c.drawString(60, y, f"共进行 {len(chat_history)} 次对话")
    y -= 20

    c.setFont(FONT_NAME, 14)
    c.drawString(50, y, "三、问卷统计")
    y -= 25
    c.setFont(FONT_NAME, 11)
    surveys = survey_data.get("surveys", [])
    c.drawString(60, y, f"共收到 {len(surveys)} 份问卷")
    y -= 20

    c.setFont(FONT_NAME, 14)
    c.drawString(50, y, "四、AI 心理建议")
    y -= 25
    c.setFont(FONT_NAME, 11)

    try:
        response = call_ai(
            [{"role": "user",
              "content": f"根据用户情绪记录{len(emotion_history)}次，问卷{len(surveys)}份，给出一句简短的心理健康建议。"}],
            stream=False, max_tokens=150
        )
        suggestion = response.choices[0].message.content
        for i in range(0, len(suggestion), 60):
            line = suggestion[i:i + 60]
            c.drawString(60, y, line)
            y -= 20
            if y < 100:
                c.showPage()
                y = height - 50
    except:
        c.drawString(60, y, "保持良好作息，关注情绪变化")

    c.save()
    temp_file.close()

    with open(temp_file.name, 'rb') as f:
        content = f.read()
    os.unlink(temp_file.name)

    return StreamingResponse(
        io.BytesIO(content),
        media_type="application/pdf",
        headers={"Content-Disposition": "attachment; filename=mindmirror_report.pdf"}
    )