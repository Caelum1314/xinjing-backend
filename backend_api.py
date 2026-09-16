from fastapi import FastAPI, File, UploadFile, Form
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
import cv2
import numpy as np
from deepface import DeepFace
import uvicorn
from zhipuai import ZhipuAI
import json
import os
from datetime import datetime
import io
import tempfile
import whisper
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re
from collections import deque

app = FastAPI()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

API_KEY = "7b9c1ea933a8499990d8c7fa8b2352cd.xgWrwaq1uIU8f9m9"
client = ZhipuAI(api_key=API_KEY)

SYSTEM_PROMPT = "你是心镜，一个温暖、真诚、专业的AI心理陪伴助手。像真人朋友一样聊天，语言自然有变化。"

DATA_DIR = "data"
os.makedirs(DATA_DIR, exist_ok=True)

CHAT_FILE = os.path.join(DATA_DIR, "chat_history.json")
EMOTION_FILE = os.path.join(DATA_DIR, "emotion_history.json")
SURVEY_FILE = os.path.join(DATA_DIR, "survey_data.json")
ACHIEVEMENT_FILE = os.path.join(DATA_DIR, "achievements.json")

EMOTION_MAP = {'happy': '开心', 'sad': '悲伤', 'angry': '愤怒', 'fear': '恐惧', 'surprise': '惊讶', 'neutral': '平静', 'disgust': '厌恶'}
EMOTION_SCORE = {"开心": 100, "平静": 70, "惊讶": 60, "悲伤": 35, "恐惧": 30, "愤怒": 25, "厌恶": 20}

ACHIEVEMENTS = {
    "first_chat": {"name": "初次对话", "desc": "完成第一次对话", "icon": "💬", "condition": "chat_count>=1"},
    "chat_10": {"name": "倾诉者", "desc": "完成10次对话", "icon": "🗣️", "condition": "chat_count>=10"},
    "first_emotion": {"name": "情绪觉察", "desc": "第一次情绪识别", "icon": "😊", "condition": "emotion_count>=1"},
    "emotion_10": {"name": "情绪记录者", "desc": "记录10次情绪", "icon": "📊", "condition": "emotion_count>=10"},
    "survey_first": {"name": "问卷新手", "desc": "第一次提交问卷", "icon": "📋", "condition": "survey_count>=1"},
    "happy_streak_3": {"name": "阳光心态", "desc": "连续3天情绪积极", "icon": "☀️", "condition": "happy_streak>=3"},
}

print("正在加载语音识别模型...")
whisper_model = whisper.load_model("base")
print("语音识别模型加载完成")

user_conversations = {}

def get_user_history(user_id="default", max_length=20):
    if user_id not in user_conversations:
        user_conversations[user_id] = deque(maxlen=max_length)
    return user_conversations[user_id]

def add_to_history(user_id, role, content):
    history = get_user_history(user_id)
    history.append({"role": role, "content": content})

def clear_user_history(user_id="default"):
    if user_id in user_conversations:
        user_conversations[user_id].clear()

def load_chat_history():
    if os.path.exists(CHAT_FILE):
        with open(CHAT_FILE, 'r') as f:
            return json.load(f)
    return []

def save_chat_history(history):
    with open(CHAT_FILE, 'w') as f:
        json.dump(history[-200:], f, ensure_ascii=False, indent=2)

def add_chat_record(user_msg, ai_msg):
    history = load_chat_history()
    history.append({
        "id": len(history) + 1,
        "user": user_msg,
        "assistant": ai_msg,
        "timestamp": datetime.now().isoformat()
    })
    save_chat_history(history)

def load_emotion_history():
    if os.path.exists(EMOTION_FILE):
        with open(EMOTION_FILE, 'r') as f:
            return json.load(f)
    return []

def save_emotion_history(history):
    with open(EMOTION_FILE, 'w') as f:
        json.dump(history[-500:], f, ensure_ascii=False, indent=2)

def add_emotion_record(emotion, confidence, source="camera"):
    history = load_emotion_history()
    history.append({
        "emotion": emotion,
        "confidence": confidence,
        "source": source,
        "timestamp": datetime.now().isoformat()
    })
    save_emotion_history(history)

def load_survey_data():
    if os.path.exists(SURVEY_FILE):
        with open(SURVEY_FILE, 'r') as f:
            return json.load(f)
    return {"surveys": []}

def save_survey_data(data):
    with open(SURVEY_FILE, 'w') as f:
        json.dump(data, f, ensure_ascii=False, indent=2)

def load_achievements():
    if os.path.exists(ACHIEVEMENT_FILE):
        with open(ACHIEVEMENT_FILE, 'r') as f:
            return json.load(f)
    return {"unlocked": [], "stats": {"chat_count": 0, "emotion_count": 0, "survey_count": 0, "happy_streak": 0, "active_days": 0, "last_active": None}}

def save_achievements(data):
    with open(ACHIEVEMENT_FILE, 'w') as f:
        json.dump(data, f, ensure_ascii=False, indent=2)

def update_stats(stat_type):
    data = load_achievements()
    stats = data["stats"]
    if stat_type == "chat":
        stats["chat_count"] = stats.get("chat_count", 0) + 1
    elif stat_type == "emotion":
        stats["emotion_count"] = stats.get("emotion_count", 0) + 1
    elif stat_type == "survey":
        stats["survey_count"] = stats.get("survey_count", 0) + 1
    today = datetime.now().strftime("%Y-%m-%d")
    if stats.get("last_active") != today:
        stats["last_active"] = today
        stats["active_days"] = stats.get("active_days", 0) + 1
    save_achievements(data)
    return check_achievements(data)

def check_achievements(data):
    stats = data["stats"]
    unlocked = set(data["unlocked"])
    new = []
    for key, ach in ACHIEVEMENTS.items():
        if key in unlocked:
            continue
        if ">=" in ach["condition"]:
            cond_key, cond_val = ach["condition"].split(">=")
            if stats.get(cond_key, 0) >= int(cond_val):
                unlocked.add(key)
                new.append({"name": ach["name"], "desc": ach["desc"], "icon": ach["icon"]})
    data["unlocked"] = list(unlocked)
    save_achievements(data)
    return new

@app.post("/chat")
async def chat(message: str = Form(...), enable_search: bool = Form(False), deep_think: bool = Form(False), user_id: str = Form("default")):
    try:
        if deep_think:
            system_prompt = SYSTEM_PROMPT + " 请进行深度思考，从多个角度分析问题。"
            max_tokens = 800
            temperature = 0.85
        else:
            system_prompt = SYSTEM_PROMPT
            max_tokens = 500
            temperature = 0.8

        add_to_history(user_id, "user", message)
        history = get_user_history(user_id)
        messages = list(history)[:-1] + [{"role": "user", "content": message}]
        messages_with_system = [{"role": "system", "content": system_prompt}] + messages

        if enable_search:
            response = client.chat.completions.create(
                model="glm-4-flash",
                messages=messages_with_system,
                temperature=temperature,
                max_tokens=max_tokens,
                stream=True,
                tools=[{"type": "web_search", "web_search": {"enable": True}}]
            )
        else:
            response = client.chat.completions.create(
                model="glm-4-flash",
                messages=messages_with_system,
                temperature=temperature,
                max_tokens=max_tokens,
                stream=True
            )

        async def generate():
            full_reply = ""
            for chunk in response:
                if chunk.choices[0].delta.content:
                    content = chunk.choices[0].delta.content
                    full_reply += content
                    yield content

                if hasattr(chunk.choices[0].delta, 'tool_calls') and chunk.choices[0].delta.tool_calls:
                    for tool in chunk.choices[0].delta.tool_calls:
                        if tool.function.name == "web_search":
                            yield f"\n\n🔍 搜索结果：\n{tool.function.arguments}\n"

            add_to_history(user_id, "assistant", full_reply)
            add_chat_record(message, full_reply)
            new_achievements = update_stats("chat")
            if new_achievements:
                yield f"\n\n🎉 解锁成就：{new_achievements[0]['name']}！"

        return StreamingResponse(generate(), media_type="text/event-stream")
    except Exception as e:
        print(f"Chat error: {e}")
        return {"reply": f"遇到问题: {str(e)}"}

@app.post("/clear_history")
async def clear_history(user_id: str = Form("default")):
    clear_user_history(user_id)
    return {"status": "ok"}

@app.post("/analyze")
async def analyze(file: UploadFile = File(...)):
    contents = await file.read()
    nparr = np.frombuffer(contents, np.uint8)
    img = cv2.imdecode(nparr, cv2.IMREAD_COLOR)

    if img is None:
        return {"emotion": "未知", "confidence": 0, "analysis": "无法读取图片"}

    try:
        result = DeepFace.analyze(img, actions=['emotion'], enforce_detection=False, silent=True)
        if result and len(result) > 0:
            emotion_raw = result[0]['dominant_emotion']
            confidence = float(result[0]['emotion'][emotion_raw])
            emotion_cn = EMOTION_MAP.get(emotion_raw, emotion_raw)

            add_emotion_record(emotion_cn, confidence)
            update_stats("emotion")

            analysis_prompt = f"""用户通过图片识别出的情绪是：{emotion_cn}（置信度{confidence:.0%}）。

请根据这个情绪，给出：
1. 情绪解读：这种情绪通常意味着什么
2. 可能原因：产生这种情绪的常见原因
3. 调节建议：3个具体可行的调节方法
4. 暖心寄语：一句鼓励的话

要求：详细、温暖、专业，150-200字。"""

            ai_response = client.chat.completions.create(
                model="glm-4-flash",
                messages=[
                    {"role": "system", "content": "你是心镜，专业的AI心理陪伴助手。你的回复要温暖、专业、详细，帮助用户理解和管理情绪。"},
                    {"role": "user", "content": analysis_prompt}
                ],
                temperature=0.7,
                max_tokens=500
            )
            analysis = ai_response.choices[0].message.content

            return {"emotion": emotion_cn, "confidence": confidence, "analysis": analysis}
    except Exception as e:
        print(f"Analyze error: {e}")

    return {"emotion": "未知", "confidence": 0, "analysis": "未能识别到有效情绪，请确保图片清晰且包含人脸"}

@app.post("/record_emotion")
async def record_emotion(emotion: str = Form(...)):
    add_emotion_record(emotion, 1.0, "manual")
    update_stats("emotion")
    return {"status": "ok"}

@app.post("/speech_to_text")
async def speech_to_text(file: UploadFile = File(...)):
    try:
        temp = tempfile.NamedTemporaryFile(delete=False, suffix=".wav")
        content = await file.read()
        temp.write(content)
        temp.close()
        result = whisper_model.transcribe(temp.name, language="zh")
        os.unlink(temp.name)
        return {"text": result["text"].strip()}
    except Exception as e:
        return {"text": "", "error": str(e)}

@app.get("/get_chat_history")
async def get_chat_history():
    return {"history": load_chat_history()}

@app.post("/clear_chat_history")
async def clear_chat_history():
    if os.path.exists(CHAT_FILE):
        os.remove(CHAT_FILE)
    return {"status": "ok"}

@app.post("/generate_survey")
async def generate_survey(message: str = Form(...)):
    default_survey = {
        "title": "心理健康调查问卷",
        "questions": [
            {"text": "你最近一周的情绪状态如何？", "options": ["非常好", "比较好", "一般", "不太好", "很糟糕"]},
            {"text": "你感到压力主要来自哪里？", "options": ["学业/工作", "人际关系", "家庭", "经济", "健康"]},
            {"text": "你是否有睡眠问题？", "options": ["没有", "偶尔", "经常", "总是"]}
        ]
    }
    return {"reply": json.dumps(default_survey, ensure_ascii=False)}

@app.post("/save_survey")
async def save_survey(answers: str = Form(...)):
    data = load_survey_data()
    data["surveys"].append({
        "answers": json.loads(answers),
        "timestamp": datetime.now().isoformat()
    })
    save_survey_data(data)
    update_stats("survey")
    return {"status": "ok"}

@app.get("/get_survey_data")
async def get_survey_data():
    return {"results": load_survey_data()["surveys"]}

@app.get("/get_statistics")
async def get_statistics():
    data = load_survey_data()
    surveys = data["surveys"]
    if not surveys:
        return {"has_data": False, "count": 0, "chart_data": {}}

    q1_counts = {}
    for survey in surveys:
        answers = survey.get("answers", [])
        if answers:
            q1_counts[answers[0]] = q1_counts.get(answers[0], 0) + 1

    return {
        "has_data": True,
        "count": len(surveys),
        "chart_data": {"labels": list(q1_counts.keys()), "values": list(q1_counts.values())}
    }

@app.post("/analyze_survey_text")
async def analyze_survey_text():
    data = load_survey_data()
    surveys = data["surveys"]
    if not surveys:
        return {"analysis": "暂无问卷数据"}

    q1_counts = {}
    for survey in surveys:
        answers = survey.get("answers", [])
        if answers:
            q1_counts[answers[0]] = q1_counts.get(answers[0], 0) + 1

    try:
        response = client.chat.completions.create(
            model="glm-4-flash",
            messages=[{"role": "user", "content": f"分析问卷结果：{q1_counts}，共{len(surveys)}份，给出简短建议"}],
            max_tokens=200
        )
        return {"analysis": response.choices[0].message.content}
    except:
        return {"analysis": "分析完成"}

@app.get("/get_achievements")
async def get_achievements():
    data = load_achievements()
    unlocked = [{"name": ACHIEVEMENTS[k]["name"], "desc": ACHIEVEMENTS[k]["desc"], "icon": ACHIEVEMENTS[k]["icon"]} for k in data["unlocked"] if k in ACHIEVEMENTS]
    return {"unlocked": unlocked, "stats": data["stats"]}

@app.get("/emotion_trend")
async def emotion_trend():
    emotions = load_emotion_history()
    if len(emotions) < 3:
        return {"error": "数据不足"}
    scores = [EMOTION_SCORE.get(e.get("emotion", "平静"), 50) for e in emotions[-7:]]
    avg = sum(scores) / len(scores)
    if avg >= 65:
        trend = "positive"
        text = "积极向好"
    elif avg >= 40:
        trend = "stable"
        text = "平稳"
    else:
        trend = "negative"
        text = "需要关注"
    return {"prediction": trend, "trend_text": text, "predicted_score": round(avg, 1), "recent_scores": scores, "analysis": "保持良好作息"}

def register_chinese_font():
    font_paths = [
        "/usr/share/fonts/truetype/wqy/wqy-microhei.ttc",
        "C:/Windows/Fonts/simhei.ttf"
    ]
    for path in font_paths:
        if os.path.exists(path):
            try:
                pdfmetrics.registerFont(TTFont('ChineseFont', path))
                return 'ChineseFont'
            except:
                continue
    return 'Helvetica'

FONT_NAME = register_chinese_font()

@app.post("/generate_report")
async def generate_report(report_type: str = Form(...), content: str = Form(...)):
    if report_type == "docx":
        doc = Document()
        title = doc.add_heading('心镜 - 心理健康报告', 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph(f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        for line in content.split('\n'):
            if line.strip() == "":
                continue
            if line.startswith('# '):
                doc.add_heading(line[2:], level=1)
                continue
            elif line.startswith('## '):
                doc.add_heading(line[3:], level=2)
                continue
            parts = re.split(r'(\*\*.*?\*\*)', line)
            if parts:
                p = doc.add_paragraph()
                for part in parts:
                    if part.startswith('**') and part.endswith('**'):
                        p.add_run(part[2:-2]).bold = True
                    else:
                        p.add_run(part)
        temp = tempfile.NamedTemporaryFile(delete=False, suffix='.docx')
        doc.save(temp.name)
        temp.close()
        with open(temp.name, 'rb') as f:
            data = f.read()
        os.unlink(temp.name)
        return StreamingResponse(io.BytesIO(data), media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", headers={"Content-Disposition": "attachment; filename=mindmirror_report.docx"})
    else:
        temp = tempfile.NamedTemporaryFile(delete=False, suffix='.pdf')
        c = canvas.Canvas(temp.name, pagesize=A4)
        y = 800
        c.setFont(FONT_NAME, 20)
        c.drawString(50, y, "心镜 - 心理健康报告")
        y -= 40
        c.setFont(FONT_NAME, 10)
        c.drawString(50, y, f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        y -= 40
        for line in content.split('\n'):
            if y < 80:
                c.showPage()
                y = 800
                c.setFont(FONT_NAME, 11)
            if line.strip() == "":
                y -= 10
                continue
            if line.startswith('# '):
                c.setFont(FONT_NAME, 16)
                c.drawString(50, y, line[2:])
                y -= 25
                c.setFont(FONT_NAME, 11)
                continue
            elif line.startswith('## '):
                c.setFont(FONT_NAME, 14)
                c.drawString(50, y, line[3:])
                y -= 20
                c.setFont(FONT_NAME, 11)
                continue
            clean = re.sub(r'\*\*(.*?)\*\*', r'\1', line)
            c.drawString(50, y, clean)
            y -= 20
        c.save()
        temp.close()
        with open(temp.name, 'rb') as f:
            data = f.read()
        os.unlink(temp.name)
        return StreamingResponse(io.BytesIO(data), media_type="application/pdf", headers={"Content-Disposition": "attachment; filename=mindmirror_report.pdf"})

@app.get("/export_report")
async def export_report():
    chat_history = load_chat_history()
    emotion_history = load_emotion_history()
    survey_data = load_survey_data()
    temp = tempfile.NamedTemporaryFile(delete=False, suffix='.pdf')
    c = canvas.Canvas(temp.name, pagesize=A4)
    y = 800
    c.setFont(FONT_NAME, 18)
    c.drawString(50, y, "心镜 - 心理健康报告")
    y -= 40
    c.setFont(FONT_NAME, 10)
    c.drawString(50, y, f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    y -= 40
    c.setFont(FONT_NAME, 14)
    c.drawString(50, y, "一、情绪统计")
    y -= 25
    c.setFont(FONT_NAME, 11)
    counts = {}
    for item in emotion_history:
        emo = item.get('emotion', '未知')
        counts[emo] = counts.get(emo, 0) + 1
    if counts:
        for emo, cnt in counts.items():
            c.drawString(60, y, f"{emo}：{cnt} 次")
            y -= 20
    else:
        c.drawString(60, y, "暂无情绪记录")
        y -= 20
    y -= 15
    c.setFont(FONT_NAME, 14)
    c.drawString(50, y, "二、对话统计")
    y -= 25
    c.setFont(FONT_NAME, 11)
    c.drawString(60, y, f"共进行 {len(chat_history)} 次对话")
    y -= 20
    c.setFont(FONT_NAME, 14)
    c.drawString(50, y, "三、问卷统计")
    y -= 25
    c.setFont(FONT_NAME, 11)
    surveys = survey_data.get("surveys", [])
    c.drawString(60, y, f"共收到 {len(surveys)} 份问卷")
    y -= 20
    c.setFont(FONT_NAME, 14)
    c.drawString(50, y, "四、AI 心理建议")
    y -= 25
    c.setFont(FONT_NAME, 11)
    try:
        resp = client.chat.completions.create(
            model="glm-4-flash",
            messages=[{"role": "user", "content": f"根据用户情绪记录{len(emotion_history)}次，问卷{len(surveys)}份，给出一句简短的心理健康建议。"}],
            max_tokens=150
        )
        sug = resp.choices[0].message.content
        for i in range(0, len(sug), 60):
            c.drawString(60, y, sug[i:i+60])
            y -= 20
    except:
        c.drawString(60, y, "保持良好作息，关注情绪变化")
    c.save()
    temp.close()
    with open(temp.name, 'rb') as f:
        data = f.read()
    os.unlink(temp.name)
    return StreamingResponse(io.BytesIO(data), media_type="application/pdf", headers={"Content-Disposition": "attachment; filename=mindmirror_report.pdf"})

if __name__ == "__main__":
    uvicorn.run(app, host="0.0.0.0", port=8000)
